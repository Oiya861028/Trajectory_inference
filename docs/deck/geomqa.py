"""Geometry QA: estimate text fit, find overlaps and margin violations.

Substitutes for a visual pass where LibreOffice is unavailable. Character-width
factors are per-point estimates for Calibri/Cambria; deliberately slightly
pessimistic so a near-miss surfaces rather than hides.
"""
import sys
from pptx import Presentation
from pptx.util import Emu

EMU = 914400.0
W, H = 13.333, 7.5
MARGIN = 0.5
CHAR_W = 0.0085    # inches of advance per point of font size, average glyph
LINE_H = 1.22      # multiple of font size

prs = Presentation(sys.argv[1] if len(sys.argv) > 1 else
                   "progress_trajectory_inference.pptx")

def inches(v):
    return (v or 0) / EMU

def est_height(shape):
    tf = shape.text_frame
    w = inches(shape.width)
    if w <= 0:
        return 0.0
    total = 0.0
    for para in tf.paragraphs:
        txt = "".join(r.text for r in para.runs)
        size = max([(r.font.size.pt if r.font.size else 18) for r in para.runs] or [18])
        if not txt:
            total += size * LINE_H / 72
            continue
        # explicit newlines force breaks
        for seg in txt.split("\n"):
            per_line = max(1, int(w / (CHAR_W * size)))
            lines = max(1, -(-len(seg) // per_line))
            total += lines * size * LINE_H / 72
        sa = para.space_after.pt if para.space_after else 0
        total += sa / 72
    return total

issues = []
for idx, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        x, y = inches(sh.left), inches(sh.top)
        w, h = inches(sh.width), inches(sh.height)
        has_text = sh.has_text_frame and sh.text_frame.text.strip()
        boxes.append((sh, x, y, w, h, has_text))
        # slide bounds / margin
        if x < MARGIN - 0.01 or y < MARGIN - 0.01 or x + w > W - MARGIN + 0.01 or y + h > H - MARGIN + 0.01:
            name = sh.shape_type
            issues.append(f"slide {idx}: MARGIN  {str(name):28s} "
                          f"x={x:.2f} y={y:.2f} -> {x+w:.2f},{y+h:.2f}"
                          + (f'  "{sh.text_frame.text[:34]}"' if has_text else ""))
        # text fit
        if has_text:
            need = est_height(sh)
            if need > h + 0.06:
                issues.append(f"slide {idx}: OVERFLOW est {need:.2f}\" into {h:.2f}\" box "
                              f'at y={y:.2f}  "{sh.text_frame.text[:46]}"')
    # overlaps between text-bearing boxes (ignore a text box sitting on its card)
    txt = [b for b in boxes if b[5]]
    for i in range(len(txt)):
        for j in range(i + 1, len(txt)):
            _, x1, y1, w1, h1, _ = txt[i]
            _, x2, y2, w2, h2, _ = txt[j]
            ox = min(x1 + w1, x2 + w2) - max(x1, x2)
            oy = min(y1 + h1, y2 + h2) - max(y1, y2)
            if ox > 0.05 and oy > 0.05:
                issues.append(f"slide {idx}: OVERLAP {ox:.2f}x{oy:.2f}\"  "
                              f'"{txt[i][0].text_frame.text[:26]}" / '
                              f'"{txt[j][0].text_frame.text[:26]}"')

print(f"{len(prs.slides._sldIdLst)} slides checked")
if not issues:
    print("no geometry issues")
for i in issues:
    print(" ", i)
